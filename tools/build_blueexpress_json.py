"""
Convierte el tarifario BluExpress (.xlsx) a data/blueexpress_tariff.json.

Uso:
    python tools/build_blueexpress_json.py <ruta_al_xlsx>

Ejecutar una sola vez cuando se actualice el tarifario.
El archivo generado queda en data/blueexpress_tariff.json y se commitea al repo.

Estructura esperada del XLSX (Nov 2024):
  - Fila 1-5: encabezados / metadata
  - Fila 6+:  col[0]=región, col[1]=provincia, col[2]=comuna, col[3]=código_posta,
              col[4..N]=tarifas por tramo de peso (en CLP con IVA)
  - Fila de encabezado de pesos: la fila que contiene "0.5" o "0,5" en col[4]
"""

import json
import os
import sys

import openpyxl


BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTPUT_FILE = os.path.join(BASE_DIR, "data", "blueexpress_tariff.json")

FACTOR_VOLUMETRICO = 4000  # cm³/kg
IVA = 1.19


def find_header_row(ws):
    """Encuentra la fila con los tramos de peso (busca la celda con valor numérico ~0.5 en col 5)."""
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True), start=1):
        if row and len(row) > 4:
            val = row[4]
            if val is not None:
                try:
                    if abs(float(str(val).replace(",", ".")) - 0.5) < 0.01:
                        return i
                except (ValueError, TypeError):
                    pass
    return None


def parse_float(val):
    if val is None:
        return None
    try:
        return float(str(val).replace(",", ".").replace("$", "").replace(".", "", str(val).count(".") - 1).strip())
    except (ValueError, TypeError):
        return None


def parse_price(val):
    """Parsea precios como 1.234 o 1234 → float."""
    if val is None:
        return None
    s = str(val).strip().replace("$", "").replace(" ", "")
    # Chilean format uses dots as thousands separator
    if s.count(".") >= 1:
        # Could be thousands separator — remove dots if no decimal part looks like decimals
        parts = s.split(".")
        if all(len(p) == 3 for p in parts[1:]):
            s = s.replace(".", "")
    try:
        return float(s)
    except ValueError:
        return None


def build_tariff(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active

    header_row_idx = find_header_row(ws)
    if header_row_idx is None:
        print("ERROR: No se encontró la fila de tramos de peso.", file=sys.stderr)
        sys.exit(1)

    # Read weight bands from header row
    rows = list(ws.iter_rows(values_only=True))
    header_row = rows[header_row_idx - 1]
    weight_bands = []
    for cell in header_row[4:]:
        if cell is None:
            break
        try:
            weight_bands.append(float(str(cell).replace(",", ".")))
        except (ValueError, TypeError):
            break

    print(f"Tramos de peso encontrados: {weight_bands}")

    # Parse communes
    comunas = {}
    for row in rows[header_row_idx:]:
        if not row or row[2] is None:
            continue

        comuna_name = str(row[2]).strip().upper()
        if not comuna_name or comuna_name in ("COMUNA", ""):
            continue

        posta = str(row[3]).strip() if row[3] is not None else ""

        tarifas_bruto = []
        for cell in row[4: 4 + len(weight_bands)]:
            price = parse_price(cell)
            tarifas_bruto.append(price)

        # Convert bruto → neto (÷ IVA)
        tarifas_neto = [
            round(p / IVA, 0) if p is not None else None
            for p in tarifas_bruto
        ]

        comunas[comuna_name] = {
            "posta": posta,
            "tarifas_neto": tarifas_neto,
        }

    output = {
        "fuente": "BluExpress propuesta nov-2024",
        "factor_volumetrico": FACTOR_VOLUMETRICO,
        "iva_aplicado": IVA,
        "nota": "tarifas_neto en CLP sin IVA. Indexar por peso_facturado_kg (max(real, volumetrico)).",
        "tramos_kg": weight_bands,
        "comunas": comunas,
    }

    os.makedirs(os.path.dirname(OUTPUT_FILE), exist_ok=True)
    with open(OUTPUT_FILE, "w", encoding="utf-8") as f:
        json.dump(output, f, ensure_ascii=False, indent=2)

    print(f"✅ {OUTPUT_FILE} generado — {len(comunas)} comunas")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(f"Uso: python {sys.argv[0]} <ruta_al_xlsx>")
        sys.exit(1)
    build_tariff(sys.argv[1])
